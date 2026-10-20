"""
Folio - research studio for slideshows and one-pagers.

Run:   python Folio.py
Needs: pip install pywebview anthropic requests google-auth google-auth-oauthlib ddgs python-pptx reportlab python-docx pypdf requests
"""

import datetime as dt
import json
import os
import re
import subprocess
import sys
import threading
import time
import webbrowser

import webview

import classroom
import engine
import images
import reminders
import render
import review
import sources
import vault
from google_link import GoogleLink

# New-project choices remembered between runs, and copied into the project for re-rendering
PREF_KEYS = ("photos", "citation", "name", "class", "font", "accent", "notes", "agenda", "questions",
             "current_event", "focus", "source_level", "wiki_ok", "blocked_sites", "max_age")
INPUT_KEYS = ("prompt", "rubric", "task", "reading_level", "tone", "slide_count", "depth", "voice_sample",
              "student_thoughts")


def _public(spec):
    """The project as the AI sees it: without Folio's own bookkeeping (_opts, _image, ...)."""
    if isinstance(spec, dict):
        return {k: _public(v) for k, v in spec.items() if not k.startswith("_")}
    if isinstance(spec, list):
        return [_public(v) for v in spec]
    return spec


def resource(name):
    base = getattr(sys, "_MEIPASS", os.path.dirname(os.path.abspath(__file__)))
    return os.path.join(base, name)


APPDATA = os.path.join(os.environ.get("APPDATA") or os.path.expanduser("~/.config"), "Folio")
SETTINGS = os.path.join(APPDATA, "settings.json")
DEFAULT_OUT = os.path.join(os.path.expanduser("~"), "Documents", "Folio Projects")
CR_CACHE = os.path.join(APPDATA, "classroom_cache.json")


def load_settings():
    s = {"provider": "Gemini", "api_key": "", "gemini_key": "", "gemini_model": "gemini-3.8-flash",
         "brave_key": "", "model": "claude-sonnet-5-5",
         "out_dir": DEFAULT_OUT, "theme": "Midnight"}
    try:
        with open(SETTINGS, encoding="utf-8") as f:
            s.update(json.load(f))
    except Exception:
        pass
    plain = [k for k in vault.KEYS if s.get(k)]  # keys still sitting in settings.json (older versions)
    for k in vault.KEYS:
        if not s.get(k):
            s[k] = vault.get(k) or ""
    if plain:
        save_settings(s)  # moves them into Credential Manager
    return s


def save_settings(s):
    """API keys go to Windows Credential Manager; everything else to settings.json."""
    os.makedirs(APPDATA, exist_ok=True)
    out = dict(s)
    for k in vault.KEYS:
        if k in out and vault.put(k, out.get(k) or ""):
            out[k] = ""  # stored safely; keep it out of the plain-text file
    with open(SETTINGS, "w", encoding="utf-8") as f:
        json.dump(out, f, indent=2)


def slug(s):
    s = re.sub(r"[^\w\s-]", "", s or "project").strip()
    return re.sub(r"\s+", " ", s)[:60] or "Project"


def open_native(path):
    if sys.platform.startswith("win"):
        os.startfile(path)  # noqa
    elif sys.platform == "darwin":
        subprocess.Popen(["open", path])
    else:
        subprocess.Popen(["xdg-open", path])


def read_any(path):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".docx":
        from docx import Document
        d = Document(path)
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        return "\n".join(p for p in parts if p.strip())
    if ext == ".pdf":
        from pypdf import PdfReader
        return "\n".join((pg.extract_text() or "") for pg in PdfReader(path).pages)
    if ext == ".pptx":
        from pptx import Presentation
        out = []
        for sl in Presentation(path).slides:
            for sh in sl.shapes:
                if sh.has_text_frame:
                    out.append(sh.text_frame.text)
        return "\n".join(out)
    with open(path, encoding="utf-8", errors="ignore") as f:
        return f.read()


def render_all(spec, theme, folder, output):
    name = slug(spec.get("title"))
    files = {}
    if output in ("Slideshow", "Both") and spec.get("slides"):
        files["pptx"] = render.build_pptx(spec, theme, os.path.join(folder, f"{name}.pptx"))
    if output in ("One-pager", "Both") and spec.get("one_pager"):
        files["pdf"] = render.build_pdf(spec, theme, os.path.join(folder, f"{name} - One-pager.pdf"))
        files["docx"] = render.build_docx(spec, theme, os.path.join(folder, f"{name} - One-pager.docx"))
    return files


class Api:
    def __init__(self):
        # Keep these private: pywebview publishes every public attribute to the page and
        # recurses into objects like the window, which hangs and leaves the buttons unwired.
        self._google = GoogleLink(APPDATA)
        self._window = None
        self._cancel = threading.Event()
        self._classroom = classroom.Classroom(on_closed=lambda: self._emit("cr_closed"))
        self.busy = False

    # ---- helpers
    def _emit(self, kind, text="", extra=None, **kw):
        payload = {"kind": kind, "text": text, "extra": extra, **kw}
        try:
            self._window.evaluate_js(f"window.Folio.event({json.dumps(payload)})")
        except Exception:
            pass

    # ---- settings
    def get_settings(self):
        return load_settings()

    def save_settings(self, data):
        s = load_settings()
        s.update({k: v for k, v in data.items() if v is not None})
        if not s.get("out_dir"):
            s["out_dir"] = DEFAULT_OUT
        save_settings(s)
        return s

    def pick_folder(self):
        r = self._window.create_file_dialog(webview.FOLDER_DIALOG)
        return r[0] if r else None

    def import_file(self):
        r = self._window.create_file_dialog(
            webview.OPEN_DIALOG, allow_multiple=False,
            file_types=("Documents (*.txt;*.docx;*.pdf;*.pptx;*.md)", "All files (*.*)"))
        if not r:
            return None
        path = r[0] if isinstance(r, (list, tuple)) else r
        try:
            txt = read_any(path).strip()
            if not txt:
                return {"error": "Couldn't find any text in that file."}
            return {"text": txt[:20000], "name": os.path.basename(path)}
        except Exception as e:
            return {"error": f"Couldn't read that file: {e}"}

    # ---- generation
    def generate(self, opts):
        if self.busy:
            return {"error": "Already working on a project."}
        s = load_settings()
        opts = dict(opts)
        opts.update(api_key=s.get("api_key"), brave_key=s.get("brave_key"), model=s.get("model"),
                    provider=s.get("provider", "Gemini"), gemini_key=s.get("gemini_key"),
                    gemini_model=s.get("gemini_model"))
        s["theme"] = opts.get("theme", "Midnight")
        s["prefs"] = {k: opts.get(k) for k in PREF_KEYS}
        save_settings(s)
        self._cancel.clear()
        self.busy = True
        threading.Thread(target=self._worker, args=(opts, s), daemon=True).start()
        return {"ok": True}

    def _photos(self, spec, folder, opts, reuse=None):
        if opts.get("photos", "Some") == "None":
            for sl in spec.get("slides") or []:
                sl.pop("image_query", None)
            (spec.get("one_pager") or {}).pop("image_query", None)
            return
        try:
            images.attach_all(spec, folder, lambda k, t, e=None: self._emit(k, t, e), self._cancel, reuse,
                              opts=opts if opts.get("photo_check", True) else None)
        except Exception as e:
            self._emit("thought", f"Couldn't fetch photos ({e}); using text layouts.")

    def _fail(self, e):
        msg = str(e)
        if isinstance(e, engine.GeminiError):
            pass
        elif "authentication" in msg.lower() or "401" in msg:
            msg = "Your Anthropic API key was rejected. Check it in Settings."
        elif "credit" in msg.lower() or "billing" in msg.lower():
            msg = "Your Anthropic account is out of credit. Add credit at console.anthropic.com."
        self._emit("error", msg[:400])

    def _worker(self, opts, s):
        try:
            g = opts.get("google") or {}
            gfile = g.get("file") if g.get("mode") == "existing" else None
            if gfile:
                self._emit("thought", f"Reading “{gfile.get('name')}” from Google Drive…")
                opts["existing_text"] = self._google.read_text(gfile)
                opts["existing_kind"] = "Google Slides deck" if gfile.get("kind") == "slides" else "Google Doc"
            spec = engine.run(opts, lambda k, t, e=None: self._emit(k, t, e), self._cancel)
            when = dt.datetime.now()
            spec["_opts"] = {k: opts.get(k) for k in PREF_KEYS if k != "photos"}
            spec["_opts"].update(date=opts.get("date") or "", accessed=when.date().isoformat())
            folder = os.path.join(s.get("out_dir") or DEFAULT_OUT,
                                  f"{slug(spec.get('title'))} ({when:%Y-%m-%d %H%M})")
            os.makedirs(folder, exist_ok=True)
            self._check_sources(spec)
            self._photos(spec, folder, opts)
            spec = self._quality(spec, folder, opts, auto_fix=opts.get("auto_fix", True))
            theme = opts.get("theme", "Midnight")
            output = opts.get("output", "Both")
            files = render_all(spec, theme, folder, output)
            if g.get("mode") in ("new", "existing"):
                gout = self._deliver_google(g, spec, files)
            elif g.get("mode") == "manual":
                gout = {"manual": True}
            else:
                gout = {}
            meta = {"google": gout, "spec": spec, "theme": theme, "output": output, "created": when.isoformat(),
                    "inputs": {k: opts.get(k) for k in INPUT_KEYS}}
            cr = opts.get("classroom") or {}
            if cr.get("href", "").startswith("https://classroom.google.com/"):
                meta["classroom"] = {k: cr.get(k) for k in ("href", "title", "due", "points")}
            self._write(folder, meta)
            self._emit("done", project=self._project(folder, meta, files))
        except engine.Cancelled:
            self._emit("cancelled")
        except Exception as e:
            self._fail(e)
        finally:
            self.busy = False

    # ---- honest quality check (review.py), with one automatic fix-up round
    def _quality(self, spec, folder, opts, auto_fix=True):
        emit = lambda k, t, e=None: self._emit(k, t, e)
        inputs = {k: opts.get(k) for k in INPUT_KEYS}
        output = opts.get("output", "Both")
        try:
            rev = review.review(spec, opts, inputs, output, emit)
        except Exception as e:
            emit("thought", f"Couldn't run the quality check ({e}).")
            return spec
        req = review.fix_request(rev)
        if not (auto_fix and req) or self._cancel.is_set():
            return spec
        emit("stage", "write")
        emit("thought", f"Fixing {len(rev['fixes'])} thing(s) the check found…")
        try:
            fixed = engine.run(dict(opts, depth="Quick", revise={"spec": _public(spec), "instruction": req}),
                               emit, self._cancel)
        except engine.Cancelled:
            raise
        except Exception as e:
            emit("thought", f"Couldn't fix it automatically ({str(e)[:120]}); keeping the first version.")
            return spec
        fixed["_opts"] = spec.get("_opts") or {}
        seen = {x["url"]: x for x in (spec.get("_all_seen") or []) + (fixed.get("_all_seen") or []) if x.get("url")}
        fixed["_all_seen"] = list(seen.values())
        emit("stage", "build")
        self._check_sources(fixed)
        self._photos(fixed, folder, opts, reuse={i["query"]: i for i in images.all_images(spec)})
        try:
            review.review(fixed, opts, inputs, output, emit)
        except Exception:
            return spec
        # never accept a "fix" that's worse: the AI sometimes drops whole parts of the project
        def score(s):
            r = s.get("_review") or {}
            return (sum(c["ok"] for c in r.get("checks") or []),
                    sum(1 for g in s.get("rubric_check") or [] if g.get("status") == "met"),
                    -sum(1 for f in r.get("facts") or [] if f.get("status") == "unverified"))
        lost = (len(fixed.get("slides") or []) < len(spec.get("slides") or []) - 1
                or (spec.get("one_pager") and not fixed.get("one_pager")))
        if lost or score(fixed) < score(spec):
            emit("thought", "The fix-up came out worse, so Folio kept the first version.")
            spec["_review"]["fix_rejected"] = True
            return spec
        fixed["_review"]["auto_fixed"] = len(rev["fixes"])
        return fixed

    def recheck(self, folder):
        """Run the quality check again on an existing project (no changes to the content)."""
        if self.busy:
            return {"error": "Already working on a project."}
        try:
            meta = self._load(folder)
            s = load_settings()
            opts = dict(meta.get("inputs") or {}, **{k: s.get(k) for k in (
                "provider", "api_key", "model", "gemini_key", "gemini_model")})
            review.review(meta["spec"], opts, meta.get("inputs") or {}, meta.get("output", "Both"))
            self._write(folder, meta)
            return self._project(folder, meta)
        except Exception as e:
            return {"error": f"Couldn't check it: {str(e)[:200]}"}

    # ---- revise with chat
    def revise(self, folder, instruction):
        if self.busy:
            return {"error": "Already working on a project."}
        if not (instruction or "").strip():
            return {"error": "Type what you'd like changed."}
        locked = self._locked(folder)
        if locked:
            return {"error": f"Close “{locked}” first, then try again."}
        try:
            meta = self._load(folder)
        except Exception as e:
            return {"error": f"Couldn't open project: {e}"}
        s = load_settings()
        old = meta["spec"]
        o = old.get("_opts") or {}
        opts = dict(meta.get("inputs") or {})
        opts.update(o)
        opts.setdefault("citation", "Simple")
        opts.setdefault("current_event", False)  # older projects: only if the change asks for one
        opts.update(api_key=s.get("api_key"), brave_key=s.get("brave_key"), model=s.get("model"),
                    provider=s.get("provider", "Gemini"), gemini_key=s.get("gemini_key"),
                    gemini_model=s.get("gemini_model"), depth="Quick", output=meta.get("output", "Both"),
                    slide_count=len(old.get("slides") or []) or opts.get("slide_count", 10),
                    photos="Some" if images.all_images(old) else "None",
                    engines={"claude": s.get("provider") == "Claude", "duckduckgo": True,
                             "brave": bool(s.get("brave_key"))},
                    revise={"spec": _public(old), "instruction": instruction.strip()})
        self._cancel.clear()
        self.busy = True
        threading.Thread(target=self._revise_worker, args=(folder, meta, opts), daemon=True).start()
        return {"ok": True}

    def _revise_worker(self, folder, meta, opts):
        try:
            old = meta["spec"]
            spec = engine.run(opts, lambda k, t, e=None: self._emit(k, t, e), self._cancel)
            spec["_opts"] = old.get("_opts") or {}
            spec["_all_seen"] = old.get("_all_seen") or spec.get("_all_seen") or []
            self._check_sources(spec)
            self._photos(spec, folder, opts, reuse={i["query"]: i for i in images.all_images(old)})
            spec = self._quality(spec, folder, opts, auto_fix=False)  # re-check, but leave the student's change as is
            hist = os.path.join(folder, "history")
            os.makedirs(hist, exist_ok=True)
            with open(os.path.join(hist, f"project-{dt.datetime.now():%Y%m%d-%H%M%S}.json"), "w",
                      encoding="utf-8") as f:
                json.dump(meta, f, indent=2)
            for p in self._files(folder).values():
                os.remove(p)
            meta = dict(meta, spec=spec, revisions=(meta.get("revisions") or []) + [opts["revise"]["instruction"]])
            if (meta.get("google") or {}).get("manual"):
                meta["google"] = {"manual": True}
            else:
                meta["google"] = {}  # Drive copies are from before the change
            files = render_all(spec, meta.get("theme", "Midnight"), folder, meta.get("output", "Both"))
            self._write(folder, meta)
            self._emit("done", project=self._project(folder, meta, files), revised=True)
        except engine.Cancelled:
            self._emit("cancelled")
        except Exception as e:
            self._fail(e)
        finally:
            self.busy = False

    def undo(self, folder):
        """Go back to the version before the last revision."""
        hist = os.path.join(folder, "history")
        snaps = sorted(f for f in os.listdir(hist) if f.endswith(".json")) if os.path.isdir(hist) else []
        if not snaps:
            return {"error": "Nothing to undo."}
        locked = self._locked(folder)
        if locked:
            return {"error": f"Close “{locked}” first, then try again."}
        try:
            with open(os.path.join(hist, snaps[-1]), encoding="utf-8") as f:
                meta = json.load(f)
            for p in self._files(folder).values():
                os.remove(p)
            files = render_all(meta["spec"], meta.get("theme", "Midnight"), folder, meta.get("output", "Both"))
            self._write(folder, meta)
            os.remove(os.path.join(hist, snaps[-1]))
            return self._project(folder, meta, files)
        except Exception as e:
            return {"error": f"Couldn't undo: {e}"}

    # ---- slide editor
    def _rerender(self, folder, meta, snapshot=True):
        """Save a history snapshot (for Undo), rebuild the files and save the project."""
        locked = self._locked(folder)
        if locked:
            raise RuntimeError(f"Close “{locked}” first, then try again.")
        if snapshot:
            hist = os.path.join(folder, "history")
            os.makedirs(hist, exist_ok=True)
            with open(os.path.join(folder, "project.json"), encoding="utf-8") as f:
                old = f.read()
            with open(os.path.join(hist, f"project-{dt.datetime.now():%Y%m%d-%H%M%S-%f}.json"), "w",
                      encoding="utf-8") as f:
                f.write(old)
        for p in self._files(folder).values():
            os.remove(p)
        files = render_all(meta["spec"], meta.get("theme", "Midnight"), folder, meta.get("output", "Both"))
        self._write(folder, meta)
        return self._project(folder, meta, files)

    def thumbs(self, folder):
        """Photo previews for the slide strip: {slide index: data URI}."""
        out = {}
        try:
            spec = self._load(folder)["spec"]
            for i, s in enumerate(spec.get("slides") or []):
                p = render.image_path(s, os.path.join(folder, "images"))
                if p:
                    out[str(i)] = images.thumb_data(p)
            p = render.image_path(spec.get("current_event"), os.path.join(folder, "images"))
            if p:
                out["news"] = images.thumb_data(p)
        except Exception:
            pass
        return out

    _EDITABLE = ("heading", "subheading", "body", "points", "notes", "stat", "stat_label", "quote", "attribution")

    def edit_slide(self, folder, idx, fields):
        try:
            meta = self._load(folder)
            s = meta["spec"]["slides"][int(idx)]
            for k, v in (fields or {}).items():
                if k in self._EDITABLE:
                    s[k] = [x.strip() for x in v if x.strip()] if k == "points" else (v or "").strip()
                elif k in ("left", "right"):  # first line is the column heading, the rest are its points
                    lines = [x.strip() for x in (v or "").splitlines() if x.strip()]
                    s[k] = {"heading": lines[0] if lines else "", "points": lines[1:]}
                elif k == "steps":  # "Label: text" per line
                    s[k] = [{"label": l.split(":", 1)[0].strip(), "text": l.split(":", 1)[1].strip() if ":" in l else ""}
                            for l in (v or "").splitlines() if l.strip()]
            return self._rerender(folder, meta)
        except Exception as e:
            return {"error": str(e)}

    def photo_options(self, query):
        try:
            return {"options": images.options(query)}
        except Exception as e:
            return {"error": f"Photo search failed: {e}"}

    def set_photo(self, folder, idx, choice):
        """choice: one of photo_options' results, or None to remove the photo."""
        try:
            meta = self._load(folder)
            s = meta["spec"]["slides"][int(idx)]
            if choice:
                info = images.use(choice, folder, f"slide{int(idx) + 1:02d}-{dt.datetime.now():%H%M%S}")
                info["query"] = choice.get("query") or s.get("image_query") or ""
                s["_image"], s["image_query"] = info, info["query"]
            else:
                s.pop("_image", None)
                s["image_query"] = ""
            return self._rerender(folder, meta)
        except Exception as e:
            return {"error": f"Couldn't change the photo: {e}"}

    # ---- presentation coach
    def personal_qs(self, fields):
        """Before starting: the parts of the assignment only the student can answer (their opinions/life)."""
        try:
            s = load_settings()
            opts = dict(fields or {}, **{k: s.get(k) for k in ("provider", "api_key", "model", "gemini_key", "gemini_model")})
            return {"questions": engine.personal_questions(opts)}
        except Exception as e:
            return {"questions": [], "error": str(e)[:200]}

    def quiz(self, folder, fresh=False):
        """Practice questions about the student's own project (cached in the project)."""
        try:
            meta = self._load(folder)
            if meta.get("quiz") and not fresh:
                return {"questions": meta["quiz"]}
            s = load_settings()
            opts = dict(meta.get("inputs") or {}, **{k: s.get(k) for k in (
                "provider", "api_key", "model", "gemini_key", "gemini_model")})
            qs = engine.make_quiz(opts, review.project_text(meta["spec"]))
            meta["quiz"] = qs
            self._write(folder, meta)
            return {"questions": qs}
        except Exception as e:
            return {"error": f"Couldn't make a quiz: {str(e)[:200]}"}

    # ---- home dashboard
    def home(self):
        """What the dashboard needs: recent projects and the last Classroom list Folio saw."""
        cache = {}
        try:
            with open(CR_CACHE, encoding="utf-8") as f:
                cache = json.load(f)
        except Exception:
            pass
        return {"projects": self.list_projects()[:6], "classroom": cache}

    def _locked(self, folder):
        """Name of an output file that's open in PowerPoint/Word/a PDF viewer, if any."""
        for p in self._files(folder).values():
            try:
                os.rename(p, p)
            except OSError:
                return os.path.basename(p)
        return None

    def _write(self, folder, meta):
        with open(os.path.join(folder, "project.json"), "w", encoding="utf-8") as f:
            json.dump(meta, f, indent=2)

    def _project(self, folder, meta, files=None):
        hist = os.path.join(folder, "history")
        return {"spec": meta["spec"], "theme": meta.get("theme"), "files": files or self._files(folder),
                "folder": folder, "google": meta.get("google") or {},
                "photos": len(images.all_images(meta["spec"])),
                "classroom": meta.get("classroom") or None,
                "can_undo": os.path.isdir(hist) and any(f.endswith(".json") for f in os.listdir(hist))}

    def cancel(self):
        self._cancel.set()

    # ---- Google
    def _deliver_google(self, g, spec, files):
        out = {}
        title = spec.get("title") or "Folio project"
        try:
            self._emit("thought", "Sending to Google Drive…")
            if g.get("mode") == "existing":
                f = g["file"]
                if f["kind"] == "slides":
                    n = self._google.append_slides(f["id"], spec)
                    out["slides"] = {**f, "note": f"Added {n} slides"}
                    if files.get("docx"):
                        out["doc"] = self._google.upload_convert(files["docx"], f"{title} - One-pager", "doc")
                else:
                    n = self._google.append_doc(f["id"], spec)
                    out["doc"] = {**f, "note": "Added to your doc"}
                    if files.get("pptx"):
                        out["slides"] = self._google.upload_convert(files["pptx"], title, "slides")
            else:
                if files.get("pptx"):
                    out["slides"] = self._google.upload_convert(files["pptx"], title, "slides")
                if files.get("docx"):
                    out["doc"] = self._google.upload_convert(files["docx"], f"{title} - One-pager", "doc")
        except PermissionError:
            out["error"] = "Your Google sign-in expired, so the files were saved on this PC only. Sign in again from Google Drive."
        except Exception as e:
            out["error"] = f"Saved on this PC, but sending to Google failed: {e}"
        return out

    def g_status(self):
        st = {"has_client": self._google.has_client(), "connected": False, "profile": None,
              "manual": bool(load_settings().get("g_manual"))}
        try:
            if self._google.connected():
                st["connected"] = True
                st["profile"] = self._google.me()
        except Exception:
            st["connected"] = False
        return st

    def g_import_client(self):
        r = self._window.create_file_dialog(webview.OPEN_DIALOG, file_types=("Google client file (*.json)",))
        if not r:
            return None
        try:
            self._google.import_client(r[0] if isinstance(r, (list, tuple)) else r)
            return self.g_status()
        except Exception as e:
            return {"error": str(e)}

    def g_sign_in(self):
        def work():
            try:
                prof = self._google.sign_in()
                self._emit("g_signed_in", profile=prof)
            except Exception as e:
                self._emit("g_error", f"Google sign-in didn't finish: {e}")
        threading.Thread(target=work, daemon=True).start()
        return {"ok": True}

    def g_set_manual(self, on):
        """No-sign-in mode for school accounts whose admin blocks outside apps."""
        s = load_settings()
        s["g_manual"] = bool(on)
        save_settings(s)
        return self.g_status()

    def g_manual_upload(self, folder=None):
        """Open Google Drive in the browser and the project folder in Explorer, so the
        files can be dragged straight into Drive."""
        webbrowser.open("https://drive.google.com/drive/my-drive")
        if folder and os.path.isdir(folder):
            try:
                open_native(folder)
            except Exception as e:
                return {"error": str(e)}
        return {"ok": True}

    # ---- source checking
    def _check_sources(self, spec):
        try:
            sources.check_spec(spec, lambda k, t, e=None: self._emit(k, t, e), self._cancel,
                               rules=sources.rules_from(spec.get("_opts")))
        except Exception as e:
            self._emit("thought", f"Couldn't check source links ({e}).")

    def check_sources(self, folder):
        """Re-check an existing project's links, drop dead ones and rebuild its files."""
        if self.busy:
            return {"error": "Already working on a project."}
        locked = self._locked(folder)
        if locked:
            return {"error": f"Close “{locked}” first, then try again."}
        try:
            meta = self._load(folder)
            summary = sources.check_spec(meta["spec"], rules=sources.rules_from(meta["spec"].get("_opts")))
            for p in self._files(folder).values():
                os.remove(p)
            files = render_all(meta["spec"], meta.get("theme", "Midnight"), folder, meta.get("output", "Both"))
            self._write(folder, meta)
            return dict(self._project(folder, meta, files), check=summary)
        except Exception as e:
            return {"error": f"Couldn't check sources: {e}"}

    # ---- Google Classroom (read-only; see classroom.py)
    def cr_open(self):
        pos = os.environ.get("FOLIO_WINDOW_POS")
        x = y = None
        if pos:
            x, y = (int(v) + 60 for v in pos.split(","))
        self._classroom.open(x, y)
        return {"ok": True}

    def cr_status(self):
        return {"open": self._classroom.is_open()}

    def cr_list(self):
        try:
            r = self._classroom.todo()
            if r.get("missing") or r.get("assigned"):
                try:  # remembered for the Home dashboard
                    with open(CR_CACHE, "w", encoding="utf-8") as f:
                        json.dump(dict(r, when=dt.datetime.now().isoformat(timespec="minutes")), f)
                except OSError:
                    pass
            return r
        except classroom.NotSignedIn:
            return {"error": "signin"}
        except Exception as e:
            return {"error": str(e)}

    def cr_import(self, href):
        try:
            a = self._classroom.assignment(href, lambda t: self._emit("cr_progress", t))
            return {"assignment": {k: a[k] for k in ("title", "href", "teacher", "points", "due")},
                    "docs": [{k: d[k] for k in ("name", "kind", "mine", "read")} for d in a["docs"]],
                    "fields": classroom.to_project_inputs(a)}
        except classroom.NotSignedIn:
            return {"error": "signin"}
        except Exception as e:
            return {"error": str(e)}

    # Editing the student's existing files (never new ones; never Turn in)
    def _cr_link(self, folder):
        meta = self._load(folder)
        cr = meta.get("classroom") or {}
        if not cr.get("href"):
            raise RuntimeError("This project isn't linked to a Classroom assignment.")
        if not self._classroom.is_open():
            self.cr_open()
            time.sleep(3)
        return meta, cr

    def cr_link(self, folder, item):
        """Link an existing project to a Classroom assignment (for projects made before importing)."""
        try:
            if not (item or {}).get("href", "").startswith("https://classroom.google.com/"):
                return {"error": "Pick an assignment."}
            meta = self._load(folder)
            meta["classroom"] = {"href": item["href"], "title": item.get("title"), "due": item.get("due")}
            self._write(folder, meta)
            return self._project(folder, meta)
        except Exception as e:
            return {"error": str(e)}

    def cr_docs(self, folder):
        try:
            _, cr = self._cr_link(folder)
            return {"docs": self._classroom.your_docs(cr["href"])}
        except classroom.NotSignedIn:
            return {"error": "Sign in to Classroom in the Classroom window first."}
        except Exception as e:
            return {"error": str(e)}

    def cr_plan(self, folder, url, instruction):
        """Draft what to add to one of the student's docs. Nothing is written yet."""
        try:
            meta, _ = self._cr_link(folder)
            self._emit("cr_progress", "Reading your doc…")
            text = self._classroom.read_doc(url)
            if not text:
                return {"error": "Couldn't read that doc."}
            self._emit("cr_progress", "Drafting what to add…")
            s = load_settings()
            opts = dict(meta.get("inputs") or {}, **{k: s.get(k) for k in (
                "provider", "api_key", "model", "gemini_key", "gemini_model")})
            edits = engine.plan_doc_edits(opts, url, text, meta["spec"], instruction or "")
            return {"edits": [{"after": e.get("after", ""), "replace": e.get("replace", ""), "text": e.get("text", ""), "why": e.get("why", "")}
                              for e in edits if (e.get("text") or "").strip()]}
        except classroom.NotSignedIn:
            return {"error": "Sign in to Classroom in the Classroom window first."}
        except Exception as e:
            return {"error": str(e)[:300]}

    def cr_research_fill(self, folder, url, instruction):
        """Read the student's doc, research each question fresh on the web, and draft answers to add
        under each one. Nothing is written yet — the student reviews the draft first."""
        try:
            meta, cr = self._cr_link(folder)
            self._emit("cr_progress", "Reading your doc…")
            text = self._classroom.read_doc(url)
            if not text:
                return {"error": "Couldn't read that doc."}
            s = load_settings()
            opts = dict(meta.get("inputs") or {}, **{k: s.get(k) for k in (
                "provider", "api_key", "model", "gemini_key", "gemini_model", "brave_key",
                "source_level", "wiki_ok", "blocked_sites", "max_age")})
            assignment = "\n".join(x for x in (
                cr.get("title"), (meta.get("inputs") or {}).get("prompt"),
                (meta.get("inputs") or {}).get("task"), instruction) if x)
            edits = engine.research_fill(opts, url, text, assignment,
                                         lambda k, t, e=None: self._emit("cr_progress", t))
            return {"edits": [{"after": e.get("after", ""), "replace": e.get("replace", ""), "text": e.get("text", ""), "why": e.get("why", "")}
                              for e in edits if (e.get("text") or "").strip()]}
        except classroom.NotSignedIn:
            return {"error": "Sign in to Classroom in the Classroom window first."}
        except Exception as e:
            return {"error": str(e)[:300]}

    def cr_write(self, folder, url, edits):
        """Type the reviewed edits into the student's existing Google Doc, then check they're there."""
        try:
            meta, cr = self._cr_link(folder)
            res = self._classroom.write_doc(url, edits, lambda t: self._emit("cr_progress", t))
            self._emit("cr_progress", "Checking the doc…")
            time.sleep(4)  # let Docs save
            # Docs turns straight quotes curly, so compare without them
            norm = lambda t: re.sub(r"\s+", " ", re.sub(r"[‘’“”'\"]", "", t or "")).strip()
            text = norm(self._classroom.read_doc(url))
            res["verified"] = sum(1 for e in edits if norm(e.get("text"))[:40] in text)
            cr.setdefault("edits", []).append({"doc": url, "when": dt.datetime.now().isoformat(timespec="minutes"),
                                               "written": len(res["written"])})
            meta["classroom"] = cr
            self._write(folder, meta)
            self._classroom.win.load_url(url)  # leave the doc open so the student can look it over
            return res
        except classroom.NotSignedIn:
            return {"error": "Sign in to Classroom in the Classroom window first."}
        except Exception as e:
            return {"error": str(e)[:300]}

    def cr_add_slides(self, folder, url, indices):
        """Add the chosen project slides to the end of the student's existing Google Slides deck."""
        try:
            meta, cr = self._cr_link(folder)
            spec = meta["spec"]
            picked = [spec["slides"][int(i)] for i in indices if 0 <= int(i) < len(spec.get("slides") or [])]
            if not picked:
                return {"error": "Pick at least one slide."}
            slides = [{"heading": s.get("heading", ""), "body": s.get("body") or s.get("quote") or "",
                       "notes": s.get("notes", ""),
                       "points": (s.get("points") or [])
                       + [f"{x.get('label')}: {x.get('text')}" for x in s.get("steps") or []]
                       + [p for side in ("left", "right") for p in
                          ([(s.get(side) or {}).get("heading")] if (s.get(side) or {}).get("heading") else [])
                          + ((s.get(side) or {}).get("points") or [])]
                       + ([f"{s.get('stat')} — {s.get('stat_label')}"] if s.get("stat") else [])}
                      for s in picked]
            res = self._classroom.add_slides(url, slides, lambda t: self._emit("cr_progress", t))
            self._emit("cr_progress", "Checking the deck…")
            time.sleep(4)
            norm = lambda t: re.sub(r"\s+", " ", re.sub(r"[‘’“”'\"]", "", t or "")).strip().lower()
            deck = norm(self._classroom.read_doc(url))
            res["verified"] = sum(1 for s in slides if s["heading"] and norm(s["heading"])[:30] in deck)
            cr.setdefault("edits", []).append({"doc": url, "when": dt.datetime.now().isoformat(timespec="minutes"),
                                               "written": res["added"], "kind": "slides"})
            meta["classroom"] = cr
            self._write(folder, meta)
            return res
        except classroom.NotSignedIn:
            return {"error": "Sign in to Classroom in the Classroom window first."}
        except Exception as e:
            return {"error": str(e)[:300]}

    def cr_show(self, href):
        try:
            self._classroom.show(href)
        except Exception as e:
            return {"error": str(e)}

    def g_sign_out(self):
        self._google.sign_out()
        return self.g_status()

    def _g(self, fn, *a):
        try:
            return fn(*a)
        except PermissionError:
            return {"error": "signed_out"}
        except Exception as e:
            return {"error": str(e)}

    def g_list(self, kind="all", query=""):
        return self._g(self._google.list_files, kind, query)

    def g_thumb(self, link):
        return self._google.thumbnail(link)

    def g_create(self, kind, title):
        return self._g(self._google.create, kind, title or ("Untitled presentation" if kind == "slides" else "Untitled document"))

    # ---- projects
    def _load(self, folder):
        with open(os.path.join(folder, "project.json"), encoding="utf-8") as f:
            return json.load(f)

    def _files(self, folder):
        out = {}
        for fn in os.listdir(folder):
            p = os.path.join(folder, fn)
            if fn.endswith(".pptx"):
                out["pptx"] = p
            elif fn.endswith(".pdf"):
                out["pdf"] = p
            elif fn.endswith(".docx"):
                out["docx"] = p
        return out

    def list_projects(self):
        root = load_settings().get("out_dir") or DEFAULT_OUT
        items = []
        if os.path.isdir(root):
            for d in os.listdir(root):
                folder = os.path.join(root, d)
                pj = os.path.join(folder, "project.json")
                if os.path.isfile(pj):
                    try:
                        m = self._load(folder)
                        when = dt.datetime.fromisoformat(m.get("created"))
                        items.append({"folder": folder, "title": m["spec"].get("title", d),
                                      "theme": m.get("theme", "Midnight"),
                                      "kind": m.get("output", "Both"),
                                      "date": when.strftime("%b %d, %Y"), "ts": when.timestamp()})
                    except Exception:
                        pass
        return sorted(items, key=lambda x: -x["ts"])

    def load_project(self, folder):
        try:
            return self._project(folder, self._load(folder))
        except Exception as e:
            return {"error": f"Couldn't open project: {e}"}

    def restyle(self, folder, theme):
        try:
            m = self._load(folder)
            for f in self._files(folder).values():
                try:
                    os.remove(f)
                except PermissionError:
                    return {"error": "Close the open slideshow / document first, then try again."}
            files = render_all(m["spec"], theme, folder, m.get("output", "Both"))
            m["theme"] = theme
            self._write(folder, m)
            return self._project(folder, m, files)
        except Exception as e:
            return {"error": f"Couldn't restyle: {e}"}

    def open_path(self, path):
        try:
            open_native(path)
        except Exception as e:
            return {"error": str(e)}

    def open_url(self, url):
        if url and url.startswith(("http://", "https://")):
            webbrowser.open(url)


def main():
    api = Api()
    with open(resource("ui.html"), encoding="utf-8") as f:
        html = f.read()
    # FOLIO_WINDOW_POS="x,y" places the window, e.g. on a second monitor for test runs
    pos = os.environ.get("FOLIO_WINDOW_POS")
    x, y = (int(v) for v in pos.split(",")) if pos else (None, None)
    win = webview.create_window("Folio", html=html, js_api=api, width=1320, height=880, x=x, y=y,
                                min_size=(1040, 700), background_color="#F5F4F0")
    api._window = win

    def quit_all():
        # the Classroom window and the hidden doc reader would otherwise keep Folio running in the background
        for w in list(webview.windows):
            if w is not win:
                try:
                    w.destroy()
                except Exception:
                    pass
    win.events.closed += quit_all

    def reminder_loop():  # missing / due-soon notifications from the last Classroom list, while Folio is open
        time.sleep(20)
        while True:
            if load_settings().get("reminders", True):
                try:
                    reminders.check(CR_CACHE, os.path.join(APPDATA, "reminders_shown.json"))
                except Exception:
                    pass
            time.sleep(3600)
    threading.Thread(target=reminder_loop, daemon=True).start()
    # A saved browser profile keeps the Classroom sign-in between launches
    webview.start(debug=False, private_mode=False, storage_path=os.path.join(APPDATA, "browser"))


if __name__ == "__main__":
    main()
