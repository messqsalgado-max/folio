"""
Folio renderers: turn a project spec into a designed .pptx deck,
plus a one-pager as .pdf and .docx.
"""

import datetime as dt
import os
import re
from urllib.parse import urlparse

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt, Emu

# ---------------------------------------------------------------- themes
THEMES = {
    "Midnight": dict(bg="12141C", surface="1C1F2B", line="2A2E3D", text="F2F3F7", muted="9EA3B8",
                     accent="8B7CF6", accent2="5FD3A3", head="Segoe UI Semibold", body="Segoe UI",
                     dark=True),
    "Paper": dict(bg="FAF7F2", surface="F0EAE0", line="E2DACB", text="1F1C18", muted="6E675C",
                  accent="B8461B", accent2="2F5D8A", head="Georgia", body="Segoe UI", dark=False),
    "Sage": dict(bg="F3F5F1", surface="E2E8DD", line="CFD8C8", text="1C2A21", muted="58695D",
                 accent="2E6A4E", accent2="C08A1E", head="Segoe UI Semibold", body="Segoe UI",
                 dark=False),
    "Ocean": dict(bg="0B1F33", surface="122D47", line="1D3E5E", text="EAF2FA", muted="93AECB",
                  accent="38BDF8", accent2="FBBF24", head="Segoe UI Semibold", body="Segoe UI",
                  dark=True),
    "Mono": dict(bg="FFFFFF", surface="F2F2F2", line="E1E1E1", text="111111", muted="666666",
                 accent="111111", accent2="E4572E", head="Segoe UI Semibold", body="Segoe UI",
                 dark=False),
}

W, H = 13.333, 7.5


def _rgb(h):
    return RGBColor.from_string(h)


def _clean(s):
    return re.sub(r"\s+", " ", (s or "")).strip()


FONTS = {"Modern": "Segoe UI Semibold", "Classic": "Georgia", "Clean": "Arial",
         "Friendly": "Trebuchet MS", "Office": "Calibri"}


def theme_for(theme, spec):
    """Theme colors/fonts with the student's own accent color and heading font applied."""
    t = dict(THEMES.get(theme, THEMES["Midnight"]))
    o = (spec or {}).get("_opts") or {}
    acc = (o.get("accent") or "").lstrip("#").upper()
    if re.fullmatch(r"[0-9A-F]{6}", acc):
        t["accent"] = acc
    if o.get("font") in FONTS:
        t["head"] = FONTS[o["font"]]
    return t


def byline(spec):
    o = (spec or {}).get("_opts") or {}
    return " · ".join(_clean(o.get(k)) for k in ("name", "class", "date") if _clean(o.get(k)))


def image_path(part, img_dir):
    img = (part or {}).get("_image")
    if not img or not img_dir:
        return None
    p = os.path.join(img_dir, img.get("file", ""))
    return p if os.path.isfile(p) else None


def _crop(path, aspect):
    """Center-crop a photo to width/height `aspect` (cached next to the original)."""
    from PIL import Image
    out = f"{os.path.splitext(path)[0]}_c{int(aspect * 1000)}.jpg"
    if os.path.isfile(out) and os.path.getmtime(out) >= os.path.getmtime(path):
        return out
    im = Image.open(path).convert("RGB")
    w, h = im.size
    if w / h > aspect:
        nw = int(h * aspect)
        im = im.crop(((w - nw) // 2, 0, (w - nw) // 2 + nw, h))
    else:
        nh = int(w / aspect)
        top = int((h - nh) * (0.5 if aspect > 2 else 0.4))  # bias up a little: faces sit high
        im = im.crop((0, top, w, top + nh))
    im.save(out, "JPEG", quality=88)
    return out


# ---------------------------------------------------------------- citations
_MON = ["January", "February", "March", "April", "May", "June", "July", "August", "September",
        "October", "November", "December"]
_MLA_MON = ["Jan.", "Feb.", "Mar.", "Apr.", "May", "June", "July", "Aug.", "Sept.", "Oct.",
            "Nov.", "Dec."]


def _parse_date(s):
    """-> (datetime, precision 'd'/'m'/'y') or (None, None)."""
    s = _clean(s).replace(".", "")
    for fmt, prec in (("%Y-%m-%d", "d"), ("%B %d, %Y", "d"), ("%b %d, %Y", "d"), ("%d %B %Y", "d"),
                      ("%d %b %Y", "d"), ("%m/%d/%Y", "d"), ("%B %Y", "m"), ("%b %Y", "m"),
                      ("%Y-%m", "m"), ("%Y", "y")):
        try:
            return dt.datetime.strptime(s, fmt), prec
        except ValueError:
            pass
    m = re.search(r"\b(1[5-9]\d\d|20\d\d)\b", s)
    return (dt.datetime(int(m.group(1)), 1, 1), "y") if m else (None, None)


def _authors(a):
    a = _clean(a)
    if not a:
        return []
    return [p for p in re.split(r"\s*(?:,\s*and\s+|\s+and\s+|&|;)\s*", a) if p]


def _is_person(n):
    org = re.search(r"\b(of|the|for|and|inc|association|society|museum|university|institute|agency|"
                    r"department|news|staff|editors|library|center|centre|council)\b", n, re.I)
    return 2 <= len(n.split()) <= 4 and not org and not n.isupper()


def _inv(n):
    if not _is_person(n):
        return n
    w = n.split()
    return f"{w[-1]}, {' '.join(w[:-1])}"


def _initials(n):
    if not _is_person(n):
        return n
    w = n.split()
    return f"{w[-1]}, " + " ".join(x[0] + "." for x in w[:-1])


def _end(s):
    s = _clean(s)
    return s if not s or s[-1] in ".?!" else s + "."


def cite(src, style="MLA", accessed=None):
    title = _clean(src.get("title")) or src.get("url", "")
    url = src.get("url", "")
    pub = _clean(src.get("publisher")) or urlparse(url).netloc.replace("www.", "")
    auth = _authors(src.get("author"))
    d, prec = _parse_date(src.get("date") or src.get("year"))
    try:
        acc = dt.date.fromisoformat(accessed) if isinstance(accessed, str) else (accessed or dt.date.today())
    except ValueError:
        acc = dt.date.today()
    if style == "MLA":
        a = ("" if not auth else _end(_inv(auth[0])) if len(auth) == 1 else
             _end(f"{_inv(auth[0])}, and {auth[1]}") if len(auth) == 2 else _end(f"{_inv(auth[0])}, et al"))
        date = "" if not d else (f"{d.day} {_MLA_MON[d.month - 1]} {d.year}" if prec == "d" else
                                 f"{_MLA_MON[d.month - 1]} {d.year}" if prec == "m" else str(d.year))
        where = ", ".join(x for x in (pub, date, re.sub(r"^https?://", "", url)) if x)
        bits = [a, f"“{_end(title)}”", where + "." if where else "",
                f"Accessed {acc.day} {_MLA_MON[acc.month - 1]} {acc.year}."]
    elif style == "APA":
        names = [_initials(x) for x in auth[:20]]
        a = "" if not names else names[0] if len(names) == 1 else ", ".join(names[:-1]) + ", & " + names[-1]
        date = "n.d." if not d else (f"{d.year}, {_MON[d.month - 1]} {d.day}" if prec == "d" else
                                     f"{d.year}, {_MON[d.month - 1]}" if prec == "m" else str(d.year))
        if a:
            bits = [_end(a), f"({date}).", _end(title), _end(pub) if pub and pub not in a else "", url]
        else:
            bits = [_end(title), f"({date}).", _end(pub), url]
    elif style == "Chicago":
        a = ("" if not auth else _end(_inv(auth[0])) if len(auth) == 1 else
             _end(_inv(auth[0]) + ", " + ", ".join(auth[1:-1] + ["and " + auth[-1]])))
        date = "" if not d else (f"{_MON[d.month - 1]} {d.day}, {d.year}" if prec == "d" else
                                 f"{_MON[d.month - 1]} {d.year}" if prec == "m" else str(d.year))
        bits = [a, f"“{_end(title)}”", _end(pub),
                _end(date) if date else f"Accessed {_MON[acc.month - 1]} {acc.day}, {acc.year}.", _end(url)]
    else:
        yr = f" ({src['year']})" if src.get("year") else ""
        return f"{title}{yr}. {pub}. {url}"
    return " ".join(b for b in bits if b)


def cite_heading(style):
    return {"MLA": "Works Cited", "APA": "References", "Chicago": "Bibliography"}.get(style, "Sources")


def citations(spec, limit=None):
    """Formatted source list in the project's citation style (alphabetical for MLA/APA/Chicago)."""
    o = spec.get("_opts") or {}
    style = o.get("citation", "Simple")
    out = [cite(x, style, o.get("accessed")) for x in spec.get("sources") or []]
    if style in ("MLA", "APA", "Chicago"):
        out.sort(key=lambda c: re.sub(r"^[\W_]+", "", c).lower())
    return out[:limit] if limit else out


def image_credits(spec):
    from images import all_images, credit
    return [f"{_clean(i.get('title'))[:70]} — {credit(i)}. {i.get('source', '')}" for i in all_images(spec)]


class Deck:
    def __init__(self, theme, spec=None, img_dir=None):
        self.t = theme_for(theme, spec)
        self.spec = spec or {}
        self.img_dir = img_dir
        self.notes_on = (self.spec.get("_opts") or {}).get("notes", True)
        self.prs = Presentation()
        self.prs.slide_width = Inches(W)
        self.prs.slide_height = Inches(H)
        self.blank = self.prs.slide_layouts[6]
        self.n = 0

    # -- primitives
    def new(self, notes=""):
        s = self.prs.slides.add_slide(self.blank)
        f = s.background.fill
        f.solid()
        f.fore_color.rgb = _rgb(self.t["bg"])
        self.n += 1
        if notes and self.notes_on:
            s.notes_slide.notes_text_frame.text = notes
        return s

    def img(self, d):
        return image_path(d, self.img_dir)

    def photo(self, s, d, x, y, w, h):
        """Photo cropped to fill the box, with a small credit line on it."""
        path = self.img(d)
        if not path:
            return False
        from images import credit
        s.shapes.add_picture(_crop(path, w / h), Inches(x), Inches(y), Inches(w), Inches(h))
        self.shade(s, x, y + h - 0.3, w, 0.3, 0.45)
        self.text(s, x + 0.12, y + h - 0.25, w - 0.24, 0.22, ("Photo: " + credit(d["_image"]))[:110],
                  size=8, color="FFFFFF", align=PP_ALIGN.RIGHT)
        return True

    def shade(self, s, x, y, w, h, alpha, color="000000"):
        """Semi-transparent box (alpha 0-1) for text on photos."""
        from pptx.oxml.ns import qn
        r = self.rect(s, x, y, w, h, color)
        clr = r._element.spPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
        clr.append(clr.makeelement(qn("a:alpha"), {"val": str(int(alpha * 100000))}))
        return r

    def rect(self, s, x, y, w, h, color, shape=MSO_SHAPE.RECTANGLE, line=None):
        r = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        r.fill.solid()
        r.fill.fore_color.rgb = _rgb(color)
        if line:
            r.line.color.rgb = _rgb(line)
            r.line.width = Pt(1)
        else:
            r.line.fill.background()
        r.shadow.inherit = False
        if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
            r.adjustments[0] = 0.06
        return r

    def text(self, s, x, y, w, h, content, size=18, color=None, bold=False, font=None,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, italic=False, spacing=1.1,
             bullets=False, gap=8):
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = Inches(0)
        tf.margin_top = tf.margin_bottom = Inches(0)
        items = content if isinstance(content, list) else [content]
        for i, item in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            p.line_spacing = spacing
            if i:
                p.space_before = Pt(gap)
            if bullets:
                r0 = p.add_run()
                r0.text = "\u25A0  "
                r0.font.size = Pt(size * 0.55)
                r0.font.color.rgb = _rgb(self.t["accent"])
                r0.font.name = self.t["body"]
            r = p.add_run()
            r.text = _clean(item)
            f = r.font
            f.size = Pt(size)
            f.bold = bold
            f.italic = italic
            f.name = font or self.t["body"]
            f.color.rgb = _rgb(color or self.t["text"])
        return tb

    def heading(self, s, txt, sub=None, width=W - 1.6):
        self.text(s, 0.8, 0.62, width, 1.0, txt, size=self._fit(txt, 34, 26, int(48 * width / (W - 1.6))),
                  bold=True, font=self.t["head"], anchor=MSO_ANCHOR.TOP)
        self.rect(s, 0.8, 1.55, 0.9, 0.06, self.t["accent"])
        if sub:
            self.text(s, 0.8, 1.75, width, 0.5, sub, size=15, color=self.t["muted"])

    def footer(self, s, title, right=W - 0.8):
        self.text(s, 0.8, H - 0.5, min(8, right - 2.6), 0.3, title, size=10, color=self.t["muted"])
        self.text(s, right - 1.0, H - 0.5, 1.0, 0.3, f"{self.n:02d}", size=10, color=self.t["muted"],
                  align=PP_ALIGN.RIGHT)

    @staticmethod
    def _fit(txt, big, small, limit):
        return big if len(txt or "") <= limit else max(small, int(big * limit / len(txt)))

    @staticmethod
    def _bullet_size(points):
        total = sum(len(p) for p in points)
        return 22 if total < 220 else 19 if total < 360 else 16

    # -- layouts
    def title(self, d, deck_title):
        s = self.new(d.get("notes", ""))
        t = self.t
        self.rect(s, 0, 0, 0.32, H, t["accent"])
        tw = 8.6
        if self.photo(s, d, W * 0.53, 0, W * 0.47, H):
            tw = W * 0.53 - 1.6
            self.rect(s, W * 0.53, 0, 0.08, H, t["accent"])
        else:
            self.rect(s, W - 4.2, -1.2, 5.6, 5.6, t["surface"], MSO_SHAPE.OVAL)
            self.rect(s, W - 2.4, 4.4, 2.0, 2.0, t["accent2"], MSO_SHAPE.OVAL)
        head = d.get("heading") or deck_title
        self.text(s, 1.1, 2.2, tw, 2.2, head, size=self._fit(head, 50, 34, int(40 * tw / 8.6)), bold=True,
                  font=t["head"], anchor=MSO_ANCHOR.BOTTOM, spacing=1.0)
        sub = d.get("subheading") or d.get("body")
        if sub:
            self.text(s, 1.1, 4.6, tw - 0.4, 1.0, sub, size=20, color=t["muted"])
        by = byline(self.spec)
        if by:
            self.rect(s, 1.1, 5.95, 0.5, 0.04, t["accent"])
            self.text(s, 1.1, 6.1, tw - 0.4, 0.4, by, size=14, color=t["muted"])
        return s

    def section(self, d, deck_title, idx):
        s = self.new(d.get("notes", ""))
        t = self.t
        self.rect(s, 0, 0, W * 0.42, H, t["surface"])
        if self.photo(s, d, 0, 0, W * 0.42, H):
            self.shade(s, 0.6, H - 2.3, 2.2, 1.5, 0.55)
            self.text(s, 0.6, H - 2.3, 2.2, 1.5, f"{idx:02d}", size=64, bold=True, color="FFFFFF",
                      font=t["head"], anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
        else:
            self.text(s, 0.8, 2.3, 4.6, 2.2, f"{idx:02d}", size=110, bold=True, color=t["accent"],
                      font=t["head"], anchor=MSO_ANCHOR.MIDDLE)
        self.text(s, W * 0.42 + 0.8, 2.6, 6.6, 1.6, d.get("heading", ""), size=38, bold=True,
                  font=t["head"], anchor=MSO_ANCHOR.BOTTOM)
        if d.get("subheading") or d.get("body"):
            self.text(s, W * 0.42 + 0.8, 4.35, 6.4, 1.4, d.get("subheading") or d.get("body"),
                      size=17, color=t["muted"])
        return s

    def bullets(self, d, deck_title):
        s = self.new(d.get("notes", ""))
        pts = (d.get("points") or [])[:6]
        if self.img(d):
            px = W * 0.58
            self.photo(s, d, px, 0, W - px, H)
            tw = px - 1.4
            self.heading(s, d.get("heading", ""), d.get("subheading"), width=tw)
            top = 2.35 if d.get("subheading") else 2.05
            body = d.get("body")
            self.text(s, 0.8, top, tw, H - top - (2.0 if body else 0.9), pts,
                      size=min(20, self._bullet_size(pts)), bullets=True, gap=12, spacing=1.15)
            if body:
                self.rect(s, 0.8, H - 1.85, 0.06, 1.0, self.t["accent"])
                self.text(s, 1.05, H - 1.85, tw - 0.25, 1.1, body, size=13, color=self.t["muted"],
                          spacing=1.2, italic=True)
            self.footer(s, deck_title, right=px - 0.4)
            return s
        self.heading(s, d.get("heading", ""), d.get("subheading"))
        top = 2.35 if d.get("subheading") else 2.05
        has_body = bool(d.get("body"))
        width = 7.6 if has_body else W - 1.6
        self.text(s, 0.8, top, width, H - top - 0.9, pts, size=self._bullet_size(pts),
                  bullets=True, gap=14, spacing=1.15)
        if has_body:
            self.rect(s, 9.0, top, W - 9.8, 3.6, self.t["surface"], MSO_SHAPE.ROUNDED_RECTANGLE)
            self.rect(s, 9.0, top, 0.08, 3.6, self.t["accent"])
            self.text(s, 9.35, top + 0.35, W - 10.5, 3.0, d["body"], size=15,
                      color=self.t["text"], spacing=1.25)
        self.footer(s, deck_title)
        return s

    def two_column(self, d, deck_title):
        s = self.new(d.get("notes", ""))
        self.heading(s, d.get("heading", ""), d.get("subheading"))
        top = 2.4 if d.get("subheading") else 2.1
        colw = (W - 1.6 - 0.5) / 2
        for i, side in enumerate(("left", "right")):
            c = d.get(side) or {}
            x = 0.8 + i * (colw + 0.5)
            self.rect(s, x, top, colw, H - top - 0.95, self.t["surface"], MSO_SHAPE.ROUNDED_RECTANGLE)
            self.rect(s, x + 0.4, top + 0.4, 0.5, 0.06, self.t["accent"] if i == 0 else self.t["accent2"])
            self.text(s, x + 0.4, top + 0.6, colw - 0.8, 0.6, c.get("heading", ""), size=22,
                      bold=True, font=self.t["head"])
            pts = (c.get("points") or [])[:5]
            self.text(s, x + 0.4, top + 1.35, colw - 0.8, H - top - 2.6, pts,
                      size=min(18, self._bullet_size(pts)), bullets=True, gap=10)
        self.footer(s, deck_title)
        return s

    def big_stat(self, d, deck_title):
        s = self.new(d.get("notes", ""))
        t = self.t
        self.text(s, 0.8, 0.62, W - 1.6, 0.6, d.get("heading", ""), size=20, bold=True,
                  color=t["muted"], font=t["head"])
        stat = d.get("stat", "")
        self.text(s, 0.8, 1.6, 6.6, 3.0, stat, size=self._fit(stat, 140, 72, 5), bold=True,
                  color=t["accent"], font=t["head"], anchor=MSO_ANCHOR.MIDDLE, spacing=0.9)
        self.text(s, 0.8, 4.7, 6.2, 1.2, d.get("stat_label", ""), size=22, font=t["head"])
        body = d.get("body") or " ".join(d.get("points") or [])
        if self.photo(s, d, 7.9, 0, W - 7.9, H):
            self.text(s, 0.8, 5.75, 6.4, 1.0, body, size=14, color=t["muted"], spacing=1.2)
            self.footer(s, deck_title, right=7.5)
            return s
        self.rect(s, 7.9, 1.6, 0.04, 4.3, t["line"])
        self.text(s, 8.4, 1.8, W - 9.2, 4.0, body, size=18, color=t["muted"], spacing=1.3,
                  anchor=MSO_ANCHOR.MIDDLE)
        self.footer(s, deck_title)
        return s

    def quote(self, d, deck_title):
        s = self.new(d.get("notes", ""))
        t = self.t
        self.rect(s, 0, 0, W, H, t["surface"])
        ox = W * 0.36 if self.photo(s, d, 0, 0, W * 0.36, H) else 0
        self.text(s, ox + 0.9, 0.6, 2, 2, "\u201C", size=160, color=t["accent"], font="Georgia")
        q = d.get("quote") or d.get("body") or ""
        qw = W - 3.2 - ox
        self.text(s, ox + 1.6, 2.0, qw, 3.2, q, size=self._fit(q, 34, 22, int(140 * qw / (W - 3.2))),
                  italic=True, font="Georgia", anchor=MSO_ANCHOR.MIDDLE, spacing=1.2)
        if d.get("attribution"):
            self.rect(s, ox + 1.6, 5.55, 0.6, 0.05, t["accent"])
            self.text(s, ox + 2.4, 5.4, qw - 0.8, 0.4, d["attribution"], size=16, color=t["muted"])
        if d.get("heading"):
            self.text(s, ox + 1.6, 0.9, qw, 0.4, d["heading"].upper(), size=12, color=t["muted"],
                      bold=True)
        return s

    def timeline(self, d, deck_title):
        s = self.new(d.get("notes", ""))
        t = self.t
        self.heading(s, d.get("heading", ""), d.get("subheading"))
        steps = (d.get("steps") or [{"label": p, "text": ""} for p in (d.get("points") or [])])[:5]
        if not steps:
            return s
        n = len(steps)
        span = W - 1.6
        y = 3.25
        self.rect(s, 0.8, y + 0.17, span, 0.04, t["line"])
        cw = span / n
        for i, st in enumerate(steps):
            cx = 0.8 + i * cw
            self.rect(s, cx, y, 0.38, 0.38, t["accent"] if i % 2 == 0 else t["accent2"], MSO_SHAPE.OVAL)
            self.text(s, cx, y + 0.65, cw - 0.3, 0.6, st.get("label", ""), size=18, bold=True,
                      font=t["head"])
            self.text(s, cx, y + 1.25, cw - 0.35, 2.4, st.get("text", ""), size=14,
                      color=t["muted"], spacing=1.2)
        self.footer(s, deck_title)
        return s

    def closing(self, d, deck_title):
        s = self.new(d.get("notes", ""))
        t = self.t
        self.rect(s, 0, 0, W, H, t["accent"] if t["dark"] else t["text"])
        fg = "FFFFFF"
        self.text(s, 0.9, 1.0, W - 1.8, 1.4, d.get("heading", "Key takeaways"), size=40,
                  bold=True, color=fg, font=t["head"])
        pts = (d.get("points") or [])[:4]
        y = 2.7
        for i, p in enumerate(pts):
            self.text(s, 0.9, y, 0.8, 0.6, f"{i + 1:02d}", size=20, bold=True,
                      color=t["accent2"] if not t["dark"] else t["bg"])
            self.text(s, 1.8, y, W - 3, 0.8, p, size=20, color=fg)
            y += 0.95
        if d.get("body") and not pts:
            self.text(s, 0.9, 2.7, W - 2, 2, d["body"], size=22, color=fg)
        return s

    def image(self, d, deck_title):
        if not self.img(d):
            return self.bullets(d, deck_title)
        s = self.new(d.get("notes", ""))
        self.photo(s, d, 0, 0, W, H)
        self.shade(s, 0, H - 2.6, W, 2.3, 0.62)
        self.rect(s, 0.8, H - 2.3, 0.9, 0.06, self.t["accent"])
        head = d.get("heading", "")
        self.text(s, 0.8, H - 2.15, W - 1.6, 0.8, head, size=self._fit(head, 34, 24, 60),
                  bold=True, color="FFFFFF", font=self.t["head"])
        cap = d.get("body") or d.get("subheading") or " ".join(d.get("points") or [])
        if cap:
            self.text(s, 0.8, H - 1.3, W - 1.6, 0.8, cap, size=16, color="E8E8E8", spacing=1.15)
        return s

    def news(self, ce, deck_title):
        """'In the News' slide for the project's current event."""
        s = self.new()
        t = self.t
        has_img = bool(self.img(ce))
        tw = W * 0.58 - 1.4 if has_img else W - 1.6
        if has_img:
            self.photo(s, ce, W * 0.58, 0, W * 0.42, H)
        self.text(s, 0.8, 0.55, 3.0, 0.35, "IN THE NEWS", size=12, bold=True, color=t["accent"])
        when = _clean(ce.get("date"))
        if when:
            self.text(s, 3.0, 0.55, tw - 2.2, 0.35, when, size=12, color=t["muted"])
        head = ce.get("headline", "")
        self.text(s, 0.8, 1.0, tw, 1.2, head, size=self._fit(head, 30, 22, int(60 * tw / (W - 1.6))),
                  bold=True, font=t["head"])
        self.text(s, 0.8, 2.35, tw, 1.7, ce.get("summary", ""), size=15, spacing=1.2)
        self.rect(s, 0.8, 4.2, tw, 2.05, t["surface"], MSO_SHAPE.ROUNDED_RECTANGLE)
        self.rect(s, 0.8, 4.2, 0.08, 2.05, t["accent"])
        self.text(s, 1.1, 4.38, tw - 0.6, 0.3, "WHY IT MATTERS", size=11, bold=True, color=t["accent"])
        self.text(s, 1.1, 4.72, tw - 0.6, 1.45, ce.get("connection", ""), size=14, spacing=1.2)
        src = " — ".join(x for x in (_clean(ce.get("source_publisher")), _clean(ce.get("source_title"))) if x)
        if src:
            self.text(s, 0.8, 6.4, tw, 0.3, "Source: " + src, size=10, color=t["muted"])
        self.footer(s, deck_title, right=(W * 0.58 - 0.4) if has_img else W - 0.8)
        return s

    def sources(self, spec, deck_title):
        lines = citations(spec)
        style = (spec.get("_opts") or {}).get("citation", "Simple")
        for i in range(0, min(len(lines), 27), 9):
            s = self.new()
            self.heading(s, cite_heading(style) + (" (cont.)" if i else ""))
            chunk = lines[i:i + 9]
            self.text(s, 0.8, 2.0, W - 1.6, H - 2.8, chunk, size=12 if len(chunk) <= 6 else 11,
                      color=self.t["muted"], gap=7)
            self.footer(s, deck_title)
        creds = image_credits(spec)
        if creds:
            s = self.new()
            self.heading(s, "Image credits")
            self.text(s, 0.8, 2.0, W - 1.6, H - 2.8, creds[:14], size=10 if len(creds) > 8 else 11,
                      color=self.t["muted"], gap=5)
            self.footer(s, deck_title)

    def save(self, path):
        self.prs.save(path)


def build_pptx(spec, theme, path):
    d = Deck(theme, spec, os.path.join(os.path.dirname(os.path.abspath(path)), "images"))
    title = spec.get("title", "Presentation")
    slides = spec.get("slides") or []
    if not slides or slides[0].get("layout") != "title":
        slides = [{"layout": "title", "heading": title, "subheading": spec.get("subtitle", "")}] + slides
    if (spec.get("current_event") or {}).get("summary"):
        at = next((i for i, s in enumerate(slides) if s.get("layout") == "closing"), len(slides))
        slides = slides[:at] + [{"layout": "_news"}] + slides[at:]
    sec = 0
    for sl in slides:
        lay = sl.get("layout", "bullets")
        if lay == "_news":
            d.news(spec["current_event"], title)
        elif lay == "title":
            d.title(sl, title)
        elif lay == "section":
            sec += 1
            d.section(sl, title, sec)
        elif lay == "two_column":
            d.two_column(sl, title)
        elif lay == "big_stat":
            d.big_stat(sl, title)
        elif lay == "quote":
            d.quote(sl, title)
        elif lay == "timeline":
            d.timeline(sl, title)
        elif lay == "closing":
            d.closing(sl, title)
        elif lay == "image":
            d.image(sl, title)
        else:
            d.bullets(sl, title)
    d.sources(spec, title)
    d.save(path)
    return path


# ---------------------------------------------------------------- one-pager PDF
def _register_fonts():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    win = os.path.join(os.environ.get("WINDIR", r"C:\Windows"), "Fonts")
    cands = {
        "Body": ["segoeui.ttf"], "Body-Bold": ["segoeuib.ttf"], "Body-Italic": ["segoeuii.ttf"],
        "Serif": ["georgia.ttf"], "Serif-Bold": ["georgiab.ttf"],
    }
    got = {}
    for name, files in cands.items():
        for f in files:
            p = os.path.join(win, f)
            if os.path.exists(p):
                try:
                    pdfmetrics.registerFont(TTFont(name, p))
                    got[name] = name
                except Exception:
                    pass
    fb = {"Body": "Helvetica", "Body-Bold": "Helvetica-Bold", "Body-Italic": "Helvetica-Oblique",
          "Serif": "Times-Roman", "Serif-Bold": "Times-Bold"}
    return {k: got.get(k, v) for k, v in fb.items()}


def build_pdf(spec, theme, path):
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.colors import HexColor, white
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import (BaseDocTemplate, Frame, PageTemplate, Paragraph, Spacer,
                                    Table, TableStyle, KeepInFrame, Image)
    from xml.sax.saxutils import escape

    t = theme_for(theme, spec)
    F = _register_fonts()
    by = byline(spec)
    hero = image_path(spec.get("one_pager"), os.path.join(os.path.dirname(os.path.abspath(path)), "images"))
    op = spec.get("one_pager") or {}
    title = op.get("title") or spec.get("title", "")
    subtitle = op.get("subtitle") or spec.get("subtitle", "")
    ink = HexColor("#" + ("1F1C18" if t["dark"] else t["text"]))
    muted = HexColor("#" + ("5E6070" if t["dark"] else t["muted"]))
    accent = HexColor("#" + t["accent"])
    accent2 = HexColor("#" + t["accent2"])
    band = HexColor("#" + (t["bg"] if t["dark"] else t["text"]))
    soft = HexColor("#" + ("EEF0F6" if t["dark"] else t["surface"]))
    serif = t["head"] == "Georgia"
    HF = F["Serif-Bold"] if serif else F["Body-Bold"]

    PW, PH = letter
    M = 42
    HEADER = 118

    def header(c, doc):
        c.saveState()
        c.setFillColor(band)
        c.rect(0, PH - HEADER, PW, HEADER, stroke=0, fill=1)
        c.setFillColor(accent)
        c.rect(0, PH - HEADER, PW, 5, stroke=0, fill=1)
        c.setFillColor(white)
        size = 26 if len(title) < 45 else 20
        c.setFont(HF, size)
        c.drawString(M, PH - 58, title[:80])
        if subtitle:
            c.setFillColor(HexColor("#D9DCE6"))
            c.setFont(F["Body"], 11.5)
            c.drawString(M, PH - 82, subtitle[:120])
        if by:
            c.setFillColor(HexColor("#B9BDCB"))
            c.setFont(F["Body"], 9)
            c.drawString(M, PH - 102, by[:120])
        c.restoreState()

    st = {
        "intro": ParagraphStyle("intro", fontName=F["Body"], fontSize=11.2, leading=16, textColor=ink),
        "h": ParagraphStyle("h", fontName=HF, fontSize=12.5, leading=16, textColor=ink,
                            spaceBefore=2, spaceAfter=3),
        "b": ParagraphStyle("b", fontName=F["Body"], fontSize=9.6, leading=13.6, textColor=ink),
        "li": ParagraphStyle("li", fontName=F["Body"], fontSize=9.6, leading=13.4, textColor=ink,
                             leftIndent=10, bulletIndent=0, spaceBefore=1.5),
        "statv": ParagraphStyle("sv", fontName=HF, fontSize=22, leading=25, textColor=accent),
        "statl": ParagraphStyle("sl", fontName=F["Body"], fontSize=8.6, leading=11, textColor=muted),
        "callh": ParagraphStyle("ch", fontName=F["Body-Bold"], fontSize=8.5, leading=11,
                                textColor=accent, spaceAfter=3),
        "callb": ParagraphStyle("cb", fontName=F["Body"], fontSize=10, leading=14, textColor=ink),
        "src": ParagraphStyle("src", fontName=F["Body"], fontSize=7, leading=9, textColor=muted),
        "srch": ParagraphStyle("srch", fontName=F["Body-Bold"], fontSize=7.5, leading=10, textColor=muted),
    }
    E = lambda s: escape(_clean(s))
    flow = []
    if hero:
        from images import credit
        hw, hh = PW - 2 * M, 128
        flow += [Image(_crop(hero, hw / hh), width=hw, height=hh),
                 Paragraph("Photo: " + E(credit(op["_image"])), st["src"]), Spacer(1, 10)]
    if op.get("intro"):
        flow += [Paragraph(E(op["intro"]), st["intro"]), Spacer(1, 12)]

    stats = (op.get("key_stats") or [])[:4]
    if stats:
        cells = [[Paragraph(E(x.get("value")), st["statv"]), Spacer(1, 2),
                  Paragraph(E(x.get("label")), st["statl"])] for x in stats]
        cw = (PW - 2 * M) / len(cells)
        tbl = Table([cells], colWidths=[cw] * len(cells))
        tbl.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), soft),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 12), ("RIGHTPADDING", (0, 0), (-1, -1), 10),
            ("TOPPADDING", (0, 0), (-1, -1), 10), ("BOTTOMPADDING", (0, 0), (-1, -1), 10),
            ("LINEBEFORE", (1, 0), (-1, -1), 0.6, HexColor("#FFFFFF")),
        ]))
        flow += [tbl, Spacer(1, 14)]

    def section_flow(sec):
        out = [Paragraph(E(sec.get("heading")), st["h"])]
        if sec.get("body"):
            out.append(Paragraph(E(sec["body"]), st["b"]))
        for b in (sec.get("bullets") or [])[:5]:
            out.append(Paragraph(E(b), st["li"], bulletText="\u2022"))
        out.append(Spacer(1, 10))
        return out

    secs = (op.get("sections") or [])[:6]
    colw = (PW - 2 * M - 22) / 2
    left, right = [], []
    for i, sec in enumerate(secs):
        (left if i % 2 == 0 else right).extend(section_flow(sec))
    callout = op.get("callout") or {}
    if callout.get("text"):
        ct = Table([[[Paragraph(E(callout.get("label") or "Why it matters").upper(), st["callh"]),
                      Paragraph(E(callout["text"]), st["callb"])]]], colWidths=[colw])
        ct.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), soft),
            ("LINEBEFORE", (0, 0), (0, -1), 3, accent2),
            ("LEFTPADDING", (0, 0), (-1, -1), 12), ("RIGHTPADDING", (0, 0), (-1, -1), 10),
            ("TOPPADDING", (0, 0), (-1, -1), 9), ("BOTTOMPADDING", (0, 0), (-1, -1), 10),
        ]))
        (right if len(right) <= len(left) else left).extend([ct, Spacer(1, 10)])
    cols = Table([[left, right]], colWidths=[colw, colw])
    cols.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (0, -1), 22),
        ("RIGHTPADDING", (1, 0), (1, -1), 0),
    ]))
    flow.append(cols)
    ce = spec.get("current_event") or {}
    if ce.get("summary"):
        label = "IN THE NEWS" + (f" · {E(ce.get('date'))}" if ce.get("date") else "")
        body = [Paragraph(label, st["callh"]),
                Paragraph(f"<b>{E(ce.get('headline'))}</b> {E(ce.get('summary'))}", st["callb"]),
                Spacer(1, 4), Paragraph(f"<i>Why it matters:</i> {E(ce.get('connection'))}", st["callb"])]
        nt = Table([[body]], colWidths=[PW - 2 * M])
        nt.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), soft),
            ("LINEBEFORE", (0, 0), (0, -1), 3, accent),
            ("LEFTPADDING", (0, 0), (-1, -1), 12), ("RIGHTPADDING", (0, 0), (-1, -1), 10),
            ("TOPPADDING", (0, 0), (-1, -1), 9), ("BOTTOMPADDING", (0, 0), (-1, -1), 10),
        ]))
        flow += [Spacer(1, 6), nt, Spacer(1, 6)]
    if op.get("conclusion"):
        flow += [Spacer(1, 4), Paragraph(E(op["conclusion"]), st["intro"])]

    srcs = citations(spec, 10)
    if srcs:
        style = (spec.get("_opts") or {}).get("citation", "Simple")
        flow += [Spacer(1, 12), Paragraph(cite_heading(style).upper(), st["srch"])]
        for i, c in enumerate(srcs, 1):
            flow.append(Paragraph(E(c) if style in ("MLA", "APA", "Chicago") else f"{i}. {E(c)}", st["src"]))

    frame_h = PH - HEADER - 26 - M
    doc = BaseDocTemplate(path, pagesize=letter, leftMargin=M, rightMargin=M,
                          topMargin=HEADER + 26, bottomMargin=M, title=title)
    frame = Frame(M, M, PW - 2 * M, frame_h, id="f", leftPadding=0, rightPadding=0,
                  topPadding=0, bottomPadding=0)
    doc.addPageTemplates([PageTemplate(id="p", frames=[frame], onPage=header)])
    doc.build([KeepInFrame(PW - 2 * M, frame_h, flow, mode="shrink")])
    return path


# ---------------------------------------------------------------- one-pager DOCX
def build_docx(spec, theme, path):
    from docx import Document
    from docx.shared import Pt as DPt, RGBColor as DRGB, Inches as DIn
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    t = theme_for(theme, spec)
    op = spec.get("one_pager") or {}
    accent = DRGB.from_string(t["accent"])
    ink = DRGB.from_string("1F1C18" if t["dark"] else t["text"])
    muted = DRGB.from_string("6E675C")
    hfont = t["head"]
    hero = image_path(op, os.path.join(os.path.dirname(os.path.abspath(path)), "images"))

    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = DIn(0.75)
        s.top_margin = s.bottom_margin = DIn(0.6)
    base = doc.styles["Normal"]
    base.font.name = "Segoe UI"
    base.font.size = DPt(10.5)
    base.element.rPr.rFonts.set(qn("w:eastAsia"), "Segoe UI")

    def para(txt, size=10.5, bold=False, color=ink, font=None, after=6, italic=False):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = DPt(after)
        r = p.add_run(_clean(txt))
        r.font.size = DPt(size)
        r.bold = bold
        r.italic = italic
        r.font.color.rgb = color
        if font:
            r.font.name = font
        return p

    para(op.get("title") or spec.get("title", ""), 24, True, ink, hfont, 2)
    if op.get("subtitle") or spec.get("subtitle"):
        para(op.get("subtitle") or spec.get("subtitle"), 12, False, muted, after=8)
    # accent rule
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "18"), ("w:space", "1"), ("w:color", t["accent"])):
        bottom.set(qn(k), v)
    bdr.append(bottom)
    pPr.append(bdr)
    p.paragraph_format.space_after = DPt(8)
    if byline(spec):
        para(byline(spec), 9.5, color=muted, after=8)
    if hero:
        from images import credit
        doc.add_picture(_crop(hero, 7.0 / 2.2), width=DIn(7.0))
        para("Photo: " + credit(op["_image"]), 7, color=muted, after=8)

    if op.get("intro"):
        para(op["intro"], 11.5, after=10)
    stats = (op.get("key_stats") or [])[:4]
    if stats:
        tbl = doc.add_table(rows=1, cols=len(stats))
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, s in enumerate(stats):
            c = tbl.rows[0].cells[i]
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:fill"), "F1EEE8")
            c._tc.get_or_add_tcPr().append(shd)
            cp = c.paragraphs[0]
            r = cp.add_run(_clean(s.get("value")))
            r.bold = True
            r.font.size = DPt(18)
            r.font.color.rgb = accent
            r.font.name = hfont
            cp2 = c.add_paragraph()
            r2 = cp2.add_run(_clean(s.get("label")))
            r2.font.size = DPt(8.5)
            r2.font.color.rgb = muted
        doc.add_paragraph().paragraph_format.space_after = DPt(2)
    for sec in (op.get("sections") or []):
        para(sec.get("heading"), 13, True, ink, hfont, 3)
        if sec.get("body"):
            para(sec["body"], after=4)
        for b in sec.get("bullets") or []:
            bp = doc.add_paragraph(style="List Bullet")
            bp.paragraph_format.space_after = DPt(2)
            r = bp.add_run(_clean(b))
            r.font.size = DPt(10.5)
        doc.paragraphs[-1].paragraph_format.space_after = DPt(8)
    co = op.get("callout") or {}
    if co.get("text"):
        para((co.get("label") or "Why it matters").upper(), 8.5, True, accent, after=1)
        para(co["text"], 11, italic=True, after=8)
    ce = spec.get("current_event") or {}
    if ce.get("summary"):
        para("IN THE NEWS" + (f" · {_clean(ce.get('date'))}" if ce.get("date") else ""), 8.5, True, accent, after=1)
        para(ce.get("headline"), 11.5, True, after=2)
        para(ce.get("summary"), after=3)
        para("Why it matters: " + _clean(ce.get("connection")), italic=True, after=8)
    if op.get("conclusion"):
        para(op["conclusion"], 11, after=10)
    srcs = citations(spec, 12)
    if srcs:
        style = (spec.get("_opts") or {}).get("citation", "Simple")
        para(cite_heading(style), 9, True, muted, after=2)
        for i, c in enumerate(srcs, 1):
            p = para(c if style in ("MLA", "APA", "Chicago") else f"{i}. {c}", 8, color=muted, after=1)
            if style in ("MLA", "APA", "Chicago"):  # hanging indent, as the styles ask
                p.paragraph_format.left_indent = DIn(0.4)
                p.paragraph_format.first_line_indent = DIn(-0.4)
    creds = image_credits(spec)
    if creds:
        para("Image credits", 9, True, muted, after=2)
        for c in creds:
            para(c, 8, color=muted, after=1)
    doc.save(path)
    return path
